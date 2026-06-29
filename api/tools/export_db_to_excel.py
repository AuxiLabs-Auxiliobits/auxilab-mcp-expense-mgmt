"""Export the local SQLite expense DB to a single .xlsx workbook — one worksheet per table.

Why: to share a *data snapshot* between developers without copying the binary .db file
(which corrupts across machines — see the WAL/SHM incident). Each dev runs this, and the
companion `import_excel_to_db.py` loads a snapshot back into a fresh/empty DB.

Usage (from the api/ directory, with the venv active):
    python tools/export_db_to_excel.py                       # -> db_snapshots/expense_db_<date>.xlsx
    python tools/export_db_to_excel.py path/to/out.xlsx      # explicit output path
    python tools/export_db_to_excel.py --db ../expense.db    # explicit source db

The schema version (alembic_version) is included so the importer can warn on mismatch.
"""

from __future__ import annotations

import argparse
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Repo-root expense.db is the app default (config._DEFAULT_SQLITE_URL).
_DEFAULT_DB = Path(__file__).resolve().parents[2] / "expense.db"


def _tables(conn: sqlite3.Connection) -> list[str]:
    rows = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' "
        "AND name NOT LIKE 'sqlite_%' ORDER BY name"
    ).fetchall()
    return [r[0] for r in rows]


def export(db_path: Path, out_path: Path) -> None:
    if not db_path.exists():
        raise SystemExit(f"DB not found: {db_path}")
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    # Lead sheet: provenance + per-table row counts.
    meta = wb.create_sheet("_README")
    meta.append(["Expense DB snapshot"])
    meta.append(["exported_at (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    meta.append(["source_db", str(db_path)])
    meta.append([])
    meta.append(["table", "rows"])

    total = 0
    for table in _tables(conn):
        cols = [c[1] for c in conn.execute(f'PRAGMA table_info("{table}")').fetchall()]
        rows = conn.execute(f'SELECT * FROM "{table}"').fetchall()
        total += len(rows)
        meta.append([table, len(rows)])

        # Excel sheet names: max 31 chars, no []:*?/\ — our table names are safe, just clamp.
        ws = wb.create_sheet(table[:31])
        ws.append(cols)
        for r in rows:
            ws.append([r[c] for c in cols])
        # Freeze the header + light width so the file is readable when opened.
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c) + 4))

    conn.close()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"exported {len(_tables(sqlite3.connect(str(db_path))))} tables, {total} rows -> {out_path}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Export the SQLite expense DB to an .xlsx workbook.")
    ap.add_argument("out", nargs="?", help="output .xlsx path (default: db_snapshots/expense_db_<date>.xlsx)")
    ap.add_argument("--db", default=str(_DEFAULT_DB), help="source sqlite db (default: repo-root expense.db)")
    args = ap.parse_args()

    db_path = Path(args.db).resolve()
    if args.out:
        out_path = Path(args.out).resolve()
    else:
        stamp = datetime.now().strftime("%Y%m%d")
        out_path = Path(__file__).resolve().parents[2] / "db_snapshots" / f"expense_db_{stamp}.xlsx"
    export(db_path, out_path)


if __name__ == "__main__":
    main()
